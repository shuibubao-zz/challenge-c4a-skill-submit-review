#!/usr/bin/env python3
"""
wechat_doc_mapper.py — Scan a folder, inventory documents, infer purpose, map to challenges.

Usage:
    python3 scripts/wechat_doc_mapper.py <FOLDER_PATH> [--output OUTPUT_XLSX]

Outputs:
    - JSON inventory to stdout (for LLM consumption)
    - Excel file to --output path (default: /mnt/user-data/outputs/elite20_doc_map.xlsx)
"""

import argparse
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path

import yaml

# ---------------------------------------------------------------------------
# Challenge registry
# ---------------------------------------------------------------------------

SKILL_DIR = Path(__file__).resolve().parent.parent
CHALLENGES_PATH = SKILL_DIR / "references" / "challenges.yaml"


def load_challenges(path: Path = CHALLENGES_PATH) -> dict:
    with open(path, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    return data.get("challenges", {})


# ---------------------------------------------------------------------------
# File inventory
# ---------------------------------------------------------------------------

KNOWN_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".pptx", ".ppt",
    ".xlsx", ".xls", ".csv", ".tsv",
    ".tex", ".bib", ".md", ".txt", ".log",
    ".json", ".jsonl", ".yaml", ".yml",
    ".zip", ".tar", ".gz", ".tar.gz", ".rar", ".7z",
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg",
    ".py", ".js", ".html", ".css", ".sh",
    ".mp4", ".mp3", ".wav", ".m4a",
    ".rtf", ".odt", ".epub",
}


def inventory_folder(folder: Path) -> list[dict]:
    """Walk folder recursively, collect file metadata."""
    items = []
    for f in sorted(folder.rglob("*")):
        if f.is_file() and not f.name.startswith("."):
            try:
                stat = f.stat()
            except OSError:
                continue
            rel = str(f.relative_to(folder))
            ext = f.suffix.lower()
            # Handle compound extensions like .tar.gz
            if rel.endswith(".tar.gz"):
                ext = ".tar.gz"
            items.append({
                "path": rel,
                "name": f.stem,
                "ext": ext,
                "size_kb": round(stat.st_size / 1024, 1),
                "modified": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d"),
                "abs_path": str(f),
            })
    return items


# ---------------------------------------------------------------------------
# Title extraction (best-effort, no heavy deps required at scan time)
# ---------------------------------------------------------------------------

def extract_title(item: dict) -> str:
    """Try to extract a meaningful title from the file. Returns best guess."""
    ext = item["ext"]
    abs_path = item["abs_path"]
    size_kb = item["size_kb"]

    # Skip very large files
    if size_kb > 50_000:
        return f"[Large file: {size_kb:.0f} KB] {item['name']}"

    try:
        if ext == ".pdf":
            return _title_from_pdf(abs_path)
        elif ext == ".docx":
            return _title_from_docx(abs_path)
        elif ext == ".pptx":
            return _title_from_pptx(abs_path)
        elif ext == ".tex":
            return _title_from_tex(abs_path)
        elif ext in (".md", ".txt", ".log"):
            return _title_from_text(abs_path)
        elif ext in (".xlsx", ".xls", ".csv"):
            return _title_from_spreadsheet(abs_path, ext)
        elif ext == ".bib":
            return _title_from_bib(abs_path)
        elif ext in (".zip", ".tar", ".tar.gz", ".rar", ".7z"):
            return f"[Archive] {item['name']}"
        elif ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"):
            return f"[Image] {item['name']}"
        elif ext in (".mp4", ".mp3", ".wav", ".m4a"):
            return f"[Media] {item['name']}"
        else:
            return item["name"]
    except Exception as e:
        return f"{item['name']} (extraction error: {type(e).__name__})"


def _title_from_pdf(path: str) -> str:
    from pypdf import PdfReader
    r = PdfReader(path)
    # Try metadata title first
    meta = r.metadata
    if meta and meta.title and meta.title.strip():
        return meta.title.strip()[:120]
    # Fall back to first line of first page
    if r.pages:
        text = (r.pages[0].extract_text() or "")[:500]
        lines = [l.strip() for l in text.split("\n") if l.strip()]
        if lines:
            return lines[0][:120]
    return Path(path).stem


def _title_from_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        # Try core properties title
        if doc.core_properties.title:
            return doc.core_properties.title.strip()[:120]
        # Try first heading
        for para in doc.paragraphs[:10]:
            if para.style.name.startswith("Heading") and para.text.strip():
                return para.text.strip()[:120]
        # First non-empty paragraph
        for para in doc.paragraphs[:5]:
            if para.text.strip():
                return para.text.strip()[:120]
    except ImportError:
        # Fallback to pandoc
        import subprocess
        result = subprocess.run(
            ["pandoc", path, "-t", "plain"],
            capture_output=True, text=True, timeout=10
        )
        lines = [l.strip() for l in result.stdout.split("\n") if l.strip()]
        if lines:
            return lines[0][:120]
    return Path(path).stem


def _title_from_pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    if prs.slides:
        slide = list(prs.slides)[0]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                return shape.text.strip()[:120]
    return Path(path).stem


def _title_from_tex(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read(3000)
    m = re.search(r"\\title\{([^}]+)\}", text)
    if m:
        return m.group(1).strip()[:120]
    # Try \section
    m = re.search(r"\\section\{([^}]+)\}", text)
    if m:
        return m.group(1).strip()[:120]
    return Path(path).stem


def _title_from_text(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        lines = f.readlines(2000)
    for line in lines[:10]:
        s = line.strip()
        if s.startswith("#"):
            return s.lstrip("#").strip()[:120]
        if s:
            return s[:120]
    return Path(path).stem


def _title_from_spreadsheet(path: str, ext: str) -> str:
    if ext in (".xlsx", ".xls"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True)
            sheet = wb.active
            headers = []
            for row in sheet.iter_rows(max_row=1, values_only=True):
                headers = [str(c) for c in row if c is not None]
            wb.close()
            if headers:
                return f"[Sheet: {sheet.title}] {', '.join(headers[:5])}"
        except Exception:
            pass
    elif ext == ".csv":
        import csv
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            try:
                headers = next(reader)
                return f"[CSV] {', '.join(headers[:5])}"
            except StopIteration:
                pass
    return Path(path).stem


def _title_from_bib(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read(2000)
    entries = re.findall(r"@\w+\{", text)
    return f"[BibTeX] {len(entries)} entries"


# ---------------------------------------------------------------------------
# Purpose inference
# ---------------------------------------------------------------------------

PURPOSE_KEYWORDS = {
    # CN keywords → purpose phrases
    "大纲": "课程大纲 (Syllabus)",
    "syllabus": "课程大纲 (Syllabus)",
    "课件": "课件/讲义 (Lecture slides)",
    "slides": "课件/讲义 (Lecture slides)",
    "翻译": "翻译文档 (Translation)",
    "translation": "翻译文档 (Translation)",
    "method": "方法说明 (Methodology doc)",
    "workflow": "工作流程 (Workflow doc)",
    "prompt": "提示词/模板 (Prompt template)",
    "paper": "学术论文 (Academic paper)",
    "论文": "学术论文 (Academic paper)",
    "abstract": "摘要 (Abstract)",
    "摘要": "摘要 (Abstract)",
    "proof": "数学证明 (Mathematical proof)",
    "证明": "数学证明 (Mathematical proof)",
    "summary": "总结文档 (Summary)",
    "总结": "总结文档 (Summary)",
    "reflection": "反思文档 (Reflection)",
    "反思": "反思文档 (Reflection)",
    "discussion": "讨论记录 (Discussion log)",
    "讨论": "讨论记录 (Discussion log)",
    "interaction": "互动记录 (Interaction log)",
    "log": "日志/记录 (Log)",
    "assignment": "作业 (Assignment)",
    "作业": "作业 (Assignment)",
    "reading": "阅读材料 (Reading material)",
    "阅读": "阅读材料 (Reading material)",
    "reference": "参考文献 (References)",
    "参考": "参考文献 (References)",
    "quiz": "测试 (Quiz)",
    "exam": "考试 (Exam)",
    "note": "笔记 (Notes)",
    "笔记": "笔记 (Notes)",
    "plan": "计划 (Plan)",
    "计划": "计划 (Plan)",
}

EXT_PURPOSE = {
    ".tex": "LaTeX 源文件 (LaTeX source)",
    ".bib": "参考文献库 (Bibliography)",
    ".pptx": "演示文稿 (Presentation)",
    ".ppt": "演示文稿 (Presentation)",
    ".xlsx": "电子表格 (Spreadsheet)",
    ".csv": "数据表 (Data table)",
    ".zip": "压缩包 (Archive)",
    ".tar.gz": "压缩包 (Archive)",
    ".png": "图片 (Image)",
    ".jpg": "图片 (Image)",
    ".jpeg": "图片 (Image)",
    ".mp4": "视频 (Video)",
    ".mp3": "音频 (Audio)",
}


def infer_purpose(item: dict, title: str) -> str:
    """Infer document purpose from filename, title, extension, and path."""
    name_lower = (item["name"] + " " + item["path"]).lower()
    title_lower = title.lower()
    combined = name_lower + " " + title_lower

    # Check keyword matches
    for kw, purpose in PURPOSE_KEYWORDS.items():
        if kw.lower() in combined:
            return purpose

    # Fall back to extension-based purpose
    if item["ext"] in EXT_PURPOSE:
        return EXT_PURPOSE[item["ext"]]

    return "待确认 (To be classified)"


# ---------------------------------------------------------------------------
# Challenge mapping
# ---------------------------------------------------------------------------

def map_to_challenges(item: dict, title: str, purpose: str, challenges: dict) -> tuple[str, list[str]]:
    """Score document against each challenge. Returns (primary, [secondaries])."""
    combined = (
        item["path"] + " " + item["name"] + " " + title + " " + purpose + " " + item["ext"]
    ).lower()

    scores = {}
    for cid, cdata in challenges.items():
        score = 0
        # Keyword match
        for kw in cdata.get("signal_keywords", []):
            if kw.lower() in combined:
                score += 2
        # Extension match
        for ext in cdata.get("signal_extensions", []):
            if item["ext"] == ext.lower():
                score += 3
        # Path segment match
        for seg in cdata.get("signal_paths", []):
            if seg.lower() in item["path"].lower():
                score += 4
        scores[cid] = score

    # Sort by score descending
    ranked = sorted(scores.items(), key=lambda x: -x[1])
    primary = "C0"
    secondaries = []

    if ranked and ranked[0][1] > 0:
        primary = ranked[0][0]
        for cid, sc in ranked[1:]:
            if sc > 2:  # threshold for secondary
                secondaries.append(cid)

    return primary, secondaries


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

def generate_excel(records: list[dict], challenges: dict, output_path: str):
    """Generate a professional Excel file with inventory, coverage, and reference sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # ---- Sheet 1: Document Inventory ----
    ws1 = wb.active
    ws1.title = "Document Inventory"
    headers1 = [
        "#", "File Name", "Extension", "Size (KB)", "Modified",
        "Extracted Title", "Inferred Purpose", "Primary Challenge",
        "Secondary Challenge(s)", "Notes"
    ]
    header_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    cell_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    for col, h in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border

    for i, rec in enumerate(records, 1):
        row_data = [
            i, rec["path"], rec["ext"], rec["size_kb"], rec["modified"],
            rec["title"], rec["purpose"], rec["primary_challenge"],
            ", ".join(rec["secondary_challenges"]), rec.get("notes", "")
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws1.cell(row=i + 1, column=col, value=val)
            cell.font = cell_font
            cell.border = thin_border
            if col in (6, 7):
                cell.alignment = Alignment(wrap_text=True)

    # Auto-width
    for col in range(1, len(headers1) + 1):
        max_len = max(
            len(str(ws1.cell(row=r, column=col).value or ""))
            for r in range(1, len(records) + 2)
        )
        ws1.column_dimensions[get_column_letter(col)].width = min(max_len + 4, 40)
    ws1.freeze_panes = "A2"

    # ---- Sheet 2: Challenge Coverage ----
    ws2 = wb.create_sheet("Challenge Coverage")
    headers2 = ["Challenge ID", "Challenge Name", "Document Count", "Files", "Missing Artifacts"]
    for col, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border

    row_idx = 2
    # Include C0 for uncategorized
    all_cids = list(challenges.keys()) + ["C0"]
    for cid in all_cids:
        matched = [r for r in records if r["primary_challenge"] == cid]
        cname = challenges[cid]["name_cn"] if cid in challenges else "未分类 (Uncategorized)"
        files_str = "; ".join(r["path"] for r in matched) if matched else "—"

        # Gap analysis
        missing = []
        if cid in challenges:
            for artifact in challenges[cid].get("expected_artifacts", []):
                artifact_lower = artifact.lower()
                found = False
                for r in records:
                    if r["primary_challenge"] == cid or cid in r["secondary_challenges"]:
                        if any(kw in r["path"].lower() for kw in artifact_lower.split("—")[0].strip().split("/")):
                            found = True
                            break
                if not found:
                    missing.append(artifact.split("—")[0].strip())
        missing_str = "; ".join(missing) if missing else "✓ All covered"

        for col, val in enumerate([cid, cname, len(matched), files_str, missing_str], 1):
            cell = ws2.cell(row=row_idx, column=col, value=val)
            cell.font = cell_font
            cell.border = thin_border
            if col in (4, 5):
                cell.alignment = Alignment(wrap_text=True)
        row_idx += 1

    for col in range(1, len(headers2) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 30
    ws2.freeze_panes = "A2"

    # ---- Sheet 3: Challenge Reference ----
    ws3 = wb.create_sheet("Challenge Reference")
    headers3 = ["ID", "Name (CN)", "Name (EN)", "Description", "Expected Artifacts"]
    for col, h in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border

    for row_idx, (cid, cdata) in enumerate(challenges.items(), 2):
        artifacts_str = "\n".join(cdata.get("expected_artifacts", []))
        for col, val in enumerate([
            cid, cdata["name_cn"], cdata["name_en"],
            cdata["description"].strip(), artifacts_str
        ], 1):
            cell = ws3.cell(row=row_idx, column=col, value=val)
            cell.font = cell_font
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True)

    for col in range(1, len(headers3) + 1):
        ws3.column_dimensions[get_column_letter(col)].width = 35
    ws3.freeze_panes = "A2"

    # Save
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"Excel saved to: {output_path}", file=sys.stderr)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="WeChat Doc → Challenge Mapper")
    parser.add_argument("folder", help="Path to the WeChat-synced folder")
    parser.add_argument(
        "--output", default="/mnt/user-data/outputs/elite20_doc_map.xlsx",
        help="Output Excel path"
    )
    parser.add_argument(
        "--challenges", default=str(CHALLENGES_PATH),
        help="Path to challenges YAML"
    )
    args = parser.parse_args()

    folder = Path(args.folder).expanduser().resolve()
    if not folder.is_dir():
        print(f"ERROR: '{folder}' is not a directory.", file=sys.stderr)
        sys.exit(1)

    # Load challenges
    challenges = load_challenges(Path(args.challenges))
    print(f"Loaded {len(challenges)} challenges.", file=sys.stderr)

    # Inventory
    items = inventory_folder(folder)
    print(f"Found {len(items)} files.", file=sys.stderr)

    if not items:
        print(json.dumps({"files": [], "summary": "No files found."}, ensure_ascii=False))
        return

    # Process each file
    records = []
    for item in items:
        title = extract_title(item)
        purpose = infer_purpose(item, title)
        primary, secondaries = map_to_challenges(item, title, purpose, challenges)
        records.append({
            **item,
            "title": title,
            "purpose": purpose,
            "primary_challenge": primary,
            "secondary_challenges": secondaries,
        })

    # Generate Excel
    generate_excel(records, challenges, args.output)

    # Output JSON for LLM consumption
    output = {
        "folder": str(folder),
        "file_count": len(records),
        "files": [
            {k: v for k, v in r.items() if k != "abs_path"}
            for r in records
        ],
        "challenge_summary": {},
    }
    for cid in list(challenges.keys()) + ["C0"]:
        matched = [r["path"] for r in records if r["primary_challenge"] == cid]
        output["challenge_summary"][cid] = {
            "count": len(matched),
            "files": matched,
        }
    print(json.dumps(output, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
